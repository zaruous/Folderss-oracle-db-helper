"""DB Helper 테이블 정의서 양식(.xlsx) 검사 도구 — 엔진(C# ExcelTemplate·VelocityTemplate)과 같은 규칙으로 미리 확인한다.

사용법:
    pip install openpyxl
    python check_template.py 양식.xlsx [variables.json]

검사:
    - 줄 지시문 짝(#foreach·#if·#sheet ↔ #end, #elseif·#else 위치), 시트 끝까지 열린 블록
    - 지시문처럼 보이는 셀이 같은 줄의 다른 값 때문에 줄 지시문이 아닌 경우(경고)
    - #sheet가 시트의 첫 지시문이고 그 위에 내용이 없는지
    - 셀 안(인라인) 지시문의 짝
    - 참조의 변수·속성 이름(variables.json 기준), 반복 변수를 그 블록 밖에서 쓴 경우
    - 수식 안 참조가 ${...} 꼴인지(아니면 그려지지 않음)
    - 고정 시트 이름 31자·금지 글자, 반복 시트 탭 이름에 참조가 있는지
결과: 오류가 있으면 종료 코드 1. 이름이 '_'로 시작하는 시트는 결과에서 빠지므로 검사하지 않는다.
"""
import json
import os
import re
import sys
import urllib.request

from openpyxl import load_workbook

VARIABLES_URL = "https://zaruous.github.io/Folderss-oracle-db-helper/table-spec-templates/variables.json"
DIRECTIVE = re.compile(r"^#\{?(foreach|if|elseif|else|end|set|sheet)\}?")
REFERENCE = re.compile(r"(?<!\\)\$(!?)(\{?)([A-Za-z_][A-Za-z0-9_]*)((?:\(\))?(?:\.[A-Za-z_][A-Za-z0-9_]*(?:\(\))?)*)(\}?)")
INLINE = re.compile(r"(?<!\\)#\{?(foreach|if|elseif|else|end|set)\b\}?")
FOREACH_ARGS = re.compile(r"^\s*\$\{?([A-Za-z_][A-Za-z0-9_]*)\}?\s+in\s+(.+?)\s*$")
SET_ARGS = re.compile(r"^\s*\$\{?([A-Za-z_][A-Za-z0-9_]*)\}?\s*=(?!=)")
LIST_MEMBERS = {"size", "length", "count", "isempty", "empty"}


def load_schema(path):
    """variables.json: 인자 → 스크립트 옆 → 스크립트 상위(저장소 구조) → 현재 폴더 → 인터넷 순으로 찾는다."""
    here = os.path.dirname(os.path.abspath(__file__))
    for candidate in [path, os.path.join(here, "variables.json"), os.path.join(here, "..", "variables.json"), os.path.join(os.getcwd(), "variables.json")]:
        if candidate and os.path.exists(candidate):
            with open(candidate, encoding="utf-8") as f:
                return json.load(f)
    if path:
        raise SystemExit(f"variables.json을 찾을 수 없습니다: {path}")
    try:
        with urllib.request.urlopen(VARIABLES_URL, timeout=10) as r:
            return json.loads(r.read().decode("utf-8"))
    except Exception as ex:
        raise SystemExit(f"variables.json을 내려받지 못했습니다({ex}). {VARIABLES_URL} 를 받아 이 스크립트 옆에 두거나 두 번째 인자로 주세요.")


def match_paren(text, start):
    """text[start] == '(' 일 때 짝 ')' 위치(따옴표 안 무시). 없으면 -1."""
    depth = 0
    i = start
    while i < len(text):
        c = text[i]
        if c in "\"'":
            close = text.find(c, i + 1)
            if close < 0:
                return -1
            i = close
        elif c == "(":
            depth += 1
        elif c == ")":
            depth -= 1
            if depth == 0:
                return i
        i += 1
    return -1


def parse_row_directive(text):
    """셀 글자가 지시문 하나뿐이면 (이름, 인자) — 아니면 None. 엔진 ExcelTemplate.ParseDirective와 같은 판단."""
    t = text.strip()
    m = DIRECTIVE.match(t)
    if not m:
        return None
    name = m.group(1)
    rest = t[m.end():]
    braced_open = len(t) > 1 and t[1] == "{"
    braced_close = m.group(0).endswith("}")
    if name in ("else", "end"):
        return (name, None) if rest.strip() == "" and braced_open == braced_close else None
    if braced_open and not braced_close:
        return None
    stripped = rest.lstrip()
    offset = len(rest) - len(stripped)
    if not stripped.startswith("(") or match_paren(rest, offset) != len(rest) - 1:
        return None
    return name, rest[offset + 1:-1]


class Checker:
    def __init__(self, schema):
        self.schema = schema
        self.objects = schema["objects"]
        self.errors = []
        self.warnings = []

    def error(self, where, message):
        self.errors.append(f"{where}: {message}")

    def warn(self, where, message):
        self.warnings.append(f"{where}: {message}")

    # ---------- 참조 ----------

    def item_type(self, source, scope):
        """#foreach 대상 식의 항목 종류: $tables → table, $t.columns → col … 모르면 None."""
        m = REFERENCE.search(source)
        if not m:
            return None
        root = m.group(3)
        path = [p.replace("()", "") for p in m.group(4).split(".") if p]
        if not path:
            spec = self.schema["roots"].get(root)
            return spec.get("item") if isinstance(spec, dict) else None
        owner = scope.get(root) or (self.schema["roots"].get(root) if isinstance(self.schema["roots"].get(root), str) else None)
        if owner in self.objects and len(path) == 1:
            member = self.objects[owner].get(path[0])
            if member and member.get("type") == "list":
                return member.get("item")
        return None

    def check_refs(self, where, text, scope, in_formula=False):
        for m in REFERENCE.finditer(text):
            braced = m.group(2) == "{"
            if in_formula and not braced:
                continue
            if braced and m.group(5) != "}":
                continue
            root = m.group(3)
            path = [p.replace("()", "") for p in m.group(4).split(".") if p]
            self.check_ref(where, m.group(0), root, path, scope)

    def check_ref(self, where, raw, root, path, scope):
        roots = self.schema["roots"]
        if root in scope:
            kind = scope[root]
        elif root in roots:
            spec = roots[root]
            kind = spec if isinstance(spec, str) else "list:" + spec["item"]
        elif root in ("foreach", "velocityCount"):
            if "#loop" not in scope:
                self.error(where, f"{raw}: $foreach·$velocityCount는 #foreach·#sheet 안에서만 씁니다")
                return
            kind = "foreach" if root == "foreach" else "number"
        else:
            hint = ""
            for name in self.objects:
                if root.lower() in (name, name + "s"):
                    hint = f" (#foreach($x in …)로 만든 반복 변수 안에서만 쓸 수 있습니다)"
            self.error(where, f"{raw}: 알 수 없는 변수 ${root}{hint} — 결과에 글자 그대로 남습니다")
            return
        for i, member in enumerate(path):
            if kind is None or kind == "any":
                return
            if kind.startswith("list:"):
                if member.lower() in LIST_MEMBERS:
                    kind = "number"
                    continue
                self.error(where, f"{raw}: 목록에는 '{member}'가 없습니다(size·isEmpty만 — 항목은 #foreach로)")
                return
            members = self.objects.get(kind)
            if members is None:
                self.error(where, f"{raw}: '{member}' 앞의 값은 속성이 없습니다")
                return
            found = next((k for k in members if k.lower() == member.lower()), None)
            if found is None:
                close = [k for k in members if k.lower().startswith(member.lower()[:3])]
                self.error(where, f"{raw}: ${kind}에 '{member}'가 없습니다" + (f" (혹시 {', '.join(close)}?)" if close else ""))
                return
            spec = members[found]
            kind = "list:" + spec["item"] if spec.get("type") == "list" else None

    # ---------- 셀 안 지시문 ----------

    def check_inline(self, where, text, scope):
        """셀 하나의 인라인 지시문 짝과 그 안의 참조(인라인 #foreach 변수 포함)."""
        local = dict(scope)
        stack = []
        pos = 0
        for m in INLINE.finditer(text):
            name = m.group(1)
            if name in ("foreach", "if", "elseif", "set"):
                after = m.end()
                while after < len(text) and text[after] in " \t":
                    after += 1
                if after >= len(text) or text[after] != "(":
                    self.error(where, f"#{name} 뒤에 ( 가 있어야 합니다")
                    continue
                close = match_paren(text, after)
                if close < 0:
                    self.error(where, f"#{name}(의 닫는 ) 가 없습니다")
                    continue
                args = text[after + 1:close]
                if name == "foreach":
                    fm = FOREACH_ARGS.match(args)
                    if not fm:
                        self.error(where, "#foreach는 #foreach($항목 in $목록) 형식이어야 합니다")
                    else:
                        local[fm.group(1)] = self.item_type(fm.group(2), local)
                        local["#loop"] = True
                        self.check_refs(where, fm.group(2), local)
                    stack.append("foreach")
                elif name == "if":
                    stack.append("if")
                    self.check_refs(where, args, local)
                elif name == "elseif":
                    if not stack or stack[-1] != "if":
                        self.error(where, "#elseif에 맞는 #if가 없습니다")
                    self.check_refs(where, args, local)
                else:
                    sm = SET_ARGS.match(args)
                    if not sm:
                        self.error(where, "#set은 #set($이름 = 값) 형식이어야 합니다")
                    else:
                        local[sm.group(1)] = "any"
            elif name == "else":
                if not stack or stack[-1] != "if":
                    self.error(where, "#else에 맞는 #if가 없습니다")
            elif name == "end":
                if not stack:
                    self.error(where, "#end에 맞는 #if·#foreach가 없습니다")
                else:
                    stack.pop()
        if stack:
            self.error(where, f"#{stack[-1]}에 맞는 #end가 없습니다(셀 안)")
        # 지시문 인자 밖의 참조
        self.check_refs(where, re.sub(r"#\{?(foreach|if|elseif|set)\}?\s*\([^)]*\)", "", text), local)

    # ---------- 시트 ----------

    def check_sheet(self, ws):
        name = ws.title
        scope = {}
        stack = []  # (지시문, 위치, 그 전 scope)
        seen_content = False
        sheet_directive = None
        for row in ws.iter_rows():
            filled = [c for c in row if c.value is not None and str(c.value) != ""]
            if not filled:
                continue
            directive = None
            if len(filled) == 1 and isinstance(filled[0].value, str) and not filled[0].value.startswith("="):
                directive = parse_row_directive(filled[0].value)
            if directive is None:
                for cell in filled:
                    where = f"'{name}'!{cell.coordinate}"
                    value = cell.value
                    if not isinstance(value, str):
                        continue
                    if value.startswith("="):
                        if re.search(r"(?<!\\)\$[A-Za-z_][A-Za-z0-9_]*\.", value) and "${" not in value:
                            self.error(where, "수식 안 참조는 ${...} 꼴로 써야 그려집니다")
                        self.check_refs(where, value, scope, in_formula=True)
                        continue
                    if len(filled) > 1 and parse_row_directive(value):
                        self.warn(where, "지시문처럼 보이지만 같은 줄에 다른 값이 있어 줄 지시문이 아닙니다(셀 안에서만 동작)")
                    if "$" in value or "#" in value:
                        self.check_inline(where, value, scope)
                seen_content = True
                continue
            kind, args = directive
            where = f"'{name}'!{filled[0].coordinate}"
            if kind == "sheet":
                if sheet_directive or stack or seen_content:
                    self.error(where, "#sheet는 시트 맨 위(다른 내용보다 먼저) 한 번만 쓸 수 있습니다")
                fm = FOREACH_ARGS.match(args or "")
                if not fm:
                    self.error(where, "#sheet는 #sheet($항목 in $목록) 형식이어야 합니다")
                else:
                    self.check_refs(where, fm.group(2), scope)
                    scope = dict(scope)
                    scope[fm.group(1)] = self.item_type(fm.group(2), scope)
                    scope["#loop"] = True
                sheet_directive = where
                continue
            if kind == "foreach":
                fm = FOREACH_ARGS.match(args or "")
                stack.append(("foreach", where, scope))
                if not fm:
                    self.error(where, "#foreach는 #foreach($항목 in $목록) 형식이어야 합니다")
                else:
                    self.check_refs(where, fm.group(2), scope)
                    item = self.item_type(fm.group(2), scope)
                    if item is None:
                        self.warn(where, f"반복 대상 {fm.group(2)}의 항목 종류를 알 수 없어 그 안의 참조는 검사하지 않습니다")
                    scope = dict(scope)
                    scope[fm.group(1)] = item
                    scope["#loop"] = True
            elif kind == "if":
                self.check_refs(where, args or "", scope)
                stack.append(("if", where, scope))
            elif kind in ("elseif", "else"):
                if not stack or stack[-1][0] != "if":
                    self.error(where, f"#{kind}에 맞는 #if가 없습니다")
                elif args:
                    self.check_refs(where, args, scope)
            elif kind == "set":
                sm = SET_ARGS.match(args or "")
                if not sm:
                    self.error(where, "#set은 #set($이름 = 값) 형식이어야 합니다")
                else:
                    scope = dict(scope)
                    scope[sm.group(1)] = "any"
            elif kind == "end":
                if not stack:
                    self.error(where, "#end에 맞는 #foreach·#if가 없습니다")
                else:
                    scope = stack.pop()[2]
        for kind, where, _ in stack:
            self.error(where, f"#{kind}에 맞는 #end가 없습니다")
        # 시트 이름: 양식 파일 자체가 올바른 엑셀이어야 한다(템플릿 글자 그대로 31자 이하, 금지 글자 없음)
        if re.search(r"[\[\]:*?/\\]", name) or len(name) > 31:
            self.error(f"'{name}'", "시트(탭) 이름은 템플릿 글자 그대로 31자 이하이고 []:*?/\\ 를 쓸 수 없습니다")
        if sheet_directive:
            if "$" not in name:
                self.warn(f"'{name}'", "반복 시트(#sheet)의 탭 이름에 참조가 없어 모든 시트가 같은 이름(+ (2), (3)…)이 됩니다 — 보통 ${table.sheetName}")
            else:
                self.check_refs(f"'{name}'(탭 이름)", name, scope)
                if "sheetName" not in name:
                    self.warn(f"'{name}'", "반복 시트 탭 이름이 ${table.sheetName}이 아니면 엔진이 31자로 자르거나 (2)를 붙여 목록의 ${t.sheetName} 링크와 어긋날 수 있습니다")


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    path = sys.argv[1]
    if os.path.basename(path).startswith("~$"):
        print("엑셀 잠금 파일(~$…)입니다. 원본 .xlsx를 지정하세요.")
        return 2
    if not path.lower().endswith(".xlsx"):
        print("오류: .xlsx 파일만 양식으로 쓸 수 있습니다.")
        return 1
    checker = Checker(load_schema(sys.argv[2] if len(sys.argv) > 2 else None))
    wb = load_workbook(path)
    checked = 0
    for ws in wb.worksheets:
        if ws.title.startswith("_"):
            continue
        checker.check_sheet(ws)
        checked += 1
    if checked == 0:
        checker.error(path, "결과에 남는 시트가 없습니다(모두 '_'로 시작)")
    for line in checker.warnings:
        print("경고  " + line)
    for line in checker.errors:
        print("오류  " + line)
    print(f"검사한 시트 {checked}개 · 오류 {len(checker.errors)}개 · 경고 {len(checker.warnings)}개")
    return 1 if checker.errors else 0


if __name__ == "__main__":
    sys.exit(main())
