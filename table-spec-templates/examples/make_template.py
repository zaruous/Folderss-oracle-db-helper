"""DB Helper 테이블 정의서 양식(.xlsx)을 처음부터 만드는 예제.

사용법:
    pip install openpyxl
    python make_template.py [저장할 파일.xlsx]

기본 저장 위치는 DB Helper 사용자 양식 폴더
    %LOCALAPPDATA%\\Folderss\\plugin-data\\zaruous.folderss-db-helper\\templates\\간단형_예제.xlsx
이고, DB Helper 팝업의 [양식] 목록을 펼치면 "내 양식 — 간단형_예제"로 보인다.

만드는 시트:
    표지                 문서 정보(${doc.*}, ${db.*})
    목록                 #foreach($t in $tables) — 테이블마다 한 줄, 이름은 테이블 시트로 가는 링크
    ${table.sheetName}  #sheet($table in $tables) — 테이블마다 시트: 열 표, 외래 키, 인덱스
    _안내                결과에서 빠지는 메모 시트

규칙(지침 guide.md 3장): 지시문 셀은 그 줄에서 값이 있는 유일한 셀이어야 한다. 수식 안 참조는 ${...} 꼴.
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "맑은 고딕"
HEADER_FILL = PatternFill("solid", fgColor="3F4759")
LABEL_FILL = PatternFill("solid", fgColor="F2F3F5")
DIRECTIVE_FILL = PatternFill("solid", fgColor="FFF2CC")
LINE = Side(style="thin", color="C9CDD4")
BOX = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)


def default_path():
    base = os.environ.get("LOCALAPPDATA", ".")
    folder = os.path.join(base, "Folderss", "plugin-data", "zaruous.folderss-db-helper", "templates")
    os.makedirs(folder, exist_ok=True)
    return os.path.join(folder, "간단형_예제.xlsx")


def directive(ws, row, text, last_col):
    """지시문 줄: A열에 지시문 하나만. 같은 줄의 다른 칸은 칠하기만 한다(값은 넣지 않는다)."""
    ws.cell(row=row, column=1, value=text).font = Font(name=FONT, size=9, italic=True, color="7F6000")
    for col in range(1, last_col + 1):
        ws.cell(row=row, column=col).fill = DIRECTIVE_FILL


def header(ws, row, first_col, titles):
    for i, title in enumerate(titles):
        cell = ws.cell(row=row, column=first_col + i, value=title)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BOX


def values(ws, row, first_col, texts, center=()):
    for i, text in enumerate(texts):
        col = first_col + i
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = Font(name=FONT)
        cell.border = BOX
        cell.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if col in center else None)


def widths(ws, *sizes):
    for i, size in enumerate(sizes):
        ws.column_dimensions[chr(ord("A") + i)].width = size
    ws.sheet_view.showGridLines = False


def cover(ws):
    widths(ws, 3, 16, 60)
    title = ws.cell(row=2, column=2, value="${doc.title}")
    title.font = Font(name=FONT, size=20, bold=True, color="1F3864")
    ws.merge_cells("B2:C2")
    facts = [
        ("프로젝트", "${doc.project}"),
        ("시스템/모듈", "${doc.system}"),
        ("대상 DB", "${db.name} (${db.address})"),
        ("DBMS", "${db.version}"),
        ("스키마", "${doc.schemas}"),
        ("스키마 기준일", "${doc.date}"),
        ("작성자", "${doc.author}"),
        ("버전", "${doc.version}"),
        ("대상 테이블", "${doc.tableCount}개"),
    ]
    for i, (label, value) in enumerate(facts):
        row = 4 + i
        a = ws.cell(row=row, column=2, value=label)
        a.font = Font(name=FONT, bold=True)
        a.fill = LABEL_FILL
        a.border = BOX
        b = ws.cell(row=row, column=3, value=value)
        b.font = Font(name=FONT)
        b.border = BOX
    note = ws.cell(row=4 + len(facts) + 1, column=2, value="${doc.generatedAt} · DB Helper로 만듦")
    note.font = Font(name=FONT, size=9, color="8A8F98")


def table_list(ws):
    widths(ws, 3, 6, 12, 28, 40, 10)
    ws.cell(row=1, column=2, value="테이블 목록").font = Font(name=FONT, size=14, bold=True)
    header(ws, 3, 2, ["No", "스키마", "테이블", "설명", "행 수"])
    directive(ws, 4, "#foreach($t in $tables)", 6)
    values(ws, 5, 2, ["${foreach.count}", "${t.owner}", None, "${t.comment}", "${t.numRows}"], center=(2,))
    link = ws.cell(row=5, column=4, value='=HYPERLINK("#\'${t.sheetName}\'!A1","${t.name}")')
    link.font = Font(name=FONT, color="2F5BD1", underline="single")
    link.border = BOX
    ws.cell(row=5, column=6).number_format = "#,##0"
    directive(ws, 6, "#end", 6)
    directive(ws, 7, "#if(!$tables)", 6)
    ws.cell(row=8, column=2, value="(대상 테이블 없음)").font = Font(name=FONT, color="8A8F98")
    directive(ws, 9, "#end", 6)
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def table_sheet(ws):
    # B No · C 열 이름 · D 논리명 · E 형식 · F NULL · G 키 · H 기본값 · I 설명
    widths(ws, 3, 5, 22, 18, 20, 7, 9, 16, 36)
    directive(ws, 1, "#sheet($table in $tables)", 9)
    title = ws.cell(row=2, column=2, value="${table.name}#if($table.logicalName) — ${table.logicalName}#end")
    title.font = Font(name=FONT, size=16, bold=True, color="1F3864")
    back = ws.cell(row=2, column=9, value='=HYPERLINK("#\'목록\'!A1","← 목록")')
    back.font = Font(name=FONT, color="2F5BD1")
    back.alignment = Alignment(horizontal="right")
    ws.cell(row=3, column=2, value="${table.owner} · ${table.typeLabel} · 열 ${table.columnCount}개#if($table.numRowsText) · 약 ${table.numRowsText}행#end").font = Font(name=FONT, color="595959")

    header(ws, 5, 2, ["No", "열 이름", "논리명", "형식", "NULL", "키", "기본값", "설명"])
    directive(ws, 6, "#foreach($col in $table.columns)", 9)
    values(ws, 7, 2, ["${col.no}", "${col.name}", "${col.logicalName}", "${col.typeLabel}", "${col.nullable}", "${col.key}", "${col.default}",
                      "#if($col.description)${col.description}#elseif($col.isFk)→ ${col.fkRef}#end"], center=(2, 6, 7))
    directive(ws, 8, "#end", 9)

    ws.cell(row=10, column=2, value="외래 키").font = Font(name=FONT, bold=True)
    header(ws, 11, 2, ["No", "이름", "열", "참조 대상", "삭제 시"])
    directive(ws, 12, "#foreach($fk in $table.foreignKeys)", 9)
    values(ws, 13, 2, ["${foreach.count}", "${fk.name}", "${fk.columns}", "${fk.target}", "${fk.deleteRule}"], center=(2,))
    directive(ws, 14, "#end", 9)
    directive(ws, 15, "#if(!$table.foreignKeys)", 9)
    ws.cell(row=16, column=2, value="(없음)").font = Font(name=FONT, color="8A8F98")
    directive(ws, 17, "#end", 9)

    ws.cell(row=19, column=2, value="인덱스").font = Font(name=FONT, bold=True)
    header(ws, 20, 2, ["No", "이름", "열", "유형"])
    directive(ws, 21, "#foreach($idx in $table.indexes)", 9)
    values(ws, 22, 2, ["${foreach.count}", "${idx.name}", "${idx.columns}", "${idx.kind}"], center=(2, 5))
    directive(ws, 23, "#end", 9)
    directive(ws, 24, "#if(!$table.indexes)", 9)
    ws.cell(row=25, column=2, value="(없음)").font = Font(name=FONT, color="8A8F98")
    directive(ws, 26, "#end", 9)

    ws.freeze_panes = "A6"  # 열 표 머리(5행)까지 고정 — #sheet 줄이 빠지면 엔진이 한 줄 올려 준다
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else default_path()
    wb = Workbook()
    cover_ws = wb.active
    cover_ws.title = "표지"
    cover(cover_ws)
    table_list(wb.create_sheet("목록"))
    table_sheet(wb.create_sheet("${table.sheetName}"))
    guide = wb.create_sheet("_안내")
    guide["A1"] = "이 시트는 이름이 '_'로 시작해 정의서 결과에서 빠집니다. 문법·변수: https://zaruous.github.io/Folderss-oracle-db-helper/table-spec-templates/"
    wb.save(path)
    print(path)


if __name__ == "__main__":
    main()
